#!/usr/bin/env python3
"""Generate a structured .docx from title + plain/markdown-ish text.

Used by skill-runtime and Go fallback. Conventions:
- Title becomes document H1 (no book-title marks in heading)
- Supports Markdown (# / ## / ###), CN section markers (一、/（一）), bullets and numbered lists
- East-Asian friendly fonts and standard page margins
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

SECTION_RE = re.compile(r"^[一二三四五六七八九十百千]+[、．.]\s*")
SUBSECTION_RE = re.compile(r"^[（(][一二三四五六七八九十百千0-9]+[）)]\s*")
NUMBERED_RE = re.compile(r"^(\d+)[\.、．]\s+")
BULLET_RE = re.compile(r"^[-*•]\s+")
MD_H3 = re.compile(r"^###\s+")
MD_H2 = re.compile(r"^##\s+")
MD_H1 = re.compile(r"^#\s+")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
INLINE_BOLD_RE = re.compile(r"\*\*([^*]+)\*\*")


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.count("|") >= 2


def _is_table_separator(line: str) -> bool:
    return bool(TABLE_SEP_RE.match(line.strip()))


def _split_table_row(line: str) -> list[str]:
    s = line.strip().strip("|")
    return [cell.strip() for cell in s.split("|")]


def _add_paragraph_with_inline_md(doc, line: str, *, list_style: str | None = None) -> None:
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    if list_style:
        p = doc.add_paragraph(style=list_style)
    else:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pos = 0
    for match in INLINE_BOLD_RE.finditer(line):
        if match.start() > pos:
            run = p.add_run(line[pos:match.start()])
            _set_run_font(run, 11)
        run = p.add_run(match.group(1))
        run.bold = True
        _set_run_font(run, 11)
        pos = match.end()
    if pos < len(line):
        run = p.add_run(line[pos:])
        _set_run_font(run, 11)
    if pos == 0 and not line:
        _set_run_font(p.add_run(line), 11)


def _add_markdown_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(cols):
            text = row[ci] if ci < len(row) else ""
            text = INLINE_BOLD_RE.sub(r"\1", text)
            cell = table.rows[ri].cells[ci]
            cell.text = text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, 10)


def normalize_title(title: str) -> str:
    s = (title or "").strip()
    for ch in "《》「」『』\"'`":
        s = s.replace(ch, "")
    s = re.sub(r"(?i)^skill[_-]?docx[_-]*", "", s)
    s = re.sub(r"(?i)_?docx$", "", s)
    s = re.sub(r"(?i)\.docx$", "", s)
    s = s.replace("__", " ").replace("_", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s or "生成文档"


def _iter_lines(content: str) -> list[str]:
    text = (content or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [ln.rstrip() for ln in text.split("\n")]
    # Drop leading/trailing empties but keep internal blank lines as paragraph breaks.
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    return lines or ["（空文档）"]


def _set_run_font(run, size_pt: float = 11, east_asia: str = "微软雅黑", western: str = "Calibri") -> None:
    from docx.shared import Pt
    from docx.oxml.ns import qn

    run.font.size = Pt(size_pt)
    run.font.name = western
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    except Exception:
        pass


def _apply_style_fonts(doc) -> None:
    from docx.shared import Pt, Cm
    from docx.oxml.ns import qn

    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    styles = doc.styles
    for style_name, size in (("Normal", 11), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        try:
            style = styles[style_name]
        except KeyError:
            continue
        font = style.font
        font.name = "Calibri"
        font.size = Pt(size)
        try:
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        except Exception:
            pass


def generate_docx(path: Path, title: str, content: str) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_LINE_SPACING
        from docx.shared import Pt
    except ImportError as exc:  # pragma: no cover
        raise SystemExit(f"python-docx is required: {exc}") from exc

    doc = Document()
    _apply_style_fonts(doc)

    heading = normalize_title(title)
    h = doc.add_heading(heading, level=1)
    for run in h.runs:
        _set_run_font(run, size_pt=16)

    lines = _iter_lines(content)
    idx = 0
    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()
        if not line:
            doc.add_paragraph("")
            idx += 1
            continue

        if _is_table_row(line) and idx + 1 < len(lines) and _is_table_separator(lines[idx + 1].strip()):
            table_rows = [_split_table_row(line)]
            idx += 2
            while idx < len(lines) and _is_table_row(lines[idx].strip()):
                table_rows.append(_split_table_row(lines[idx].strip()))
                idx += 1
            _add_markdown_table(doc, table_rows)
            continue

        if MD_H3.match(line):
            p = doc.add_heading(MD_H3.sub("", line).strip(), level=3)
        elif MD_H2.match(line):
            p = doc.add_heading(MD_H2.sub("", line).strip(), level=2)
        elif MD_H1.match(line):
            text = MD_H1.sub("", line).strip()
            if text == heading:
                idx += 1
                continue
            p = doc.add_heading(text, level=1)
        elif SECTION_RE.match(line):
            p = doc.add_heading(line, level=2)
        elif SUBSECTION_RE.match(line):
            p = doc.add_heading(line, level=3)
        elif BULLET_RE.match(line):
            _add_paragraph_with_inline_md(doc, BULLET_RE.sub("", line), list_style="List Bullet")
            idx += 1
            continue
        elif NUMBERED_RE.match(line):
            _add_paragraph_with_inline_md(doc, NUMBERED_RE.sub("", line), list_style="List Number")
            idx += 1
            continue
        else:
            _add_paragraph_with_inline_md(doc, line)
            idx += 1
            continue

        for run in p.runs:
            _set_run_font(run, 12 if p.style and "Heading" in (p.style.name or "") else 11)
        idx += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _paragraph_block(paragraph) -> dict:
    text = (paragraph.text or "").strip()
    if not text:
        return {"type": "blank"}
    style = (paragraph.style.name or "").lower()
    if "heading 1" in style:
        return {"type": "h1", "text": text}
    if "heading 2" in style:
        return {"type": "h2", "text": text}
    if "heading 3" in style:
        return {"type": "h3", "text": text}
    if "list bullet" in style:
        return {"type": "li", "text": text, "ordered": False}
    if "list number" in style:
        return {"type": "li", "text": text, "ordered": True}
    return {"type": "p", "text": text}


def _table_blocks(table) -> list[dict]:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([(cell.text or "").strip() for cell in row.cells])
    if not rows:
        return []
    blocks: list[dict] = []
    for ri, row in enumerate(rows):
        blocks.append({"type": "p", "text": "| " + " | ".join(row) + " |"})
        if ri == 0:
            blocks.append({"type": "p", "text": "| " + " | ".join(["---"] * len(row)) + " |"})
    return blocks


def preview_docx(path: Path) -> dict:
    """Extract structured blocks from an existing .docx for UI preview."""
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise SystemExit(f"python-docx is required: {exc}") from exc

    doc = Document(str(path))
    blocks: list[dict] = []
    title = ""

    def append_block(block: dict) -> None:
        nonlocal title
        if block.get("type") == "blank":
            return
        if not title and block.get("type") == "h1":
            title = str(block.get("text") or "")
        blocks.append(block)

    for child in doc.element.body:
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, doc)
            block = _paragraph_block(paragraph)
            append_block(block)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            for block in _table_blocks(table):
                append_block(block)

    if not title:
        title = normalize_title(path.stem)
    return {
        "title": title,
        "filename": path.name,
        "blocks": blocks,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate a Word (.docx) artifact")
    parser.add_argument("--out", help="Output .docx path")
    parser.add_argument("--preview-json", help="Read .docx and print preview JSON to stdout")
    parser.add_argument("--title", default="生成文档")
    parser.add_argument("--content-file", help="Read body text from file (UTF-8)")
    parser.add_argument("--content", default="", help="Body text (ignored if --content-file set)")
    args = parser.parse_args()
    if args.preview_json:
        import json

        payload = preview_docx(Path(args.preview_json))
        print(json.dumps(payload, ensure_ascii=False))
        return 0
    if not args.out:
        parser.error("--out is required unless --preview-json is set")
    content = args.content
    if args.content_file:
        content = Path(args.content_file).read_text(encoding="utf-8")
    out = Path(args.out)
    generate_docx(out, args.title, content)
    print(out.resolve())
    return 0


if __name__ == "__main__":
    sys.exit(main())
